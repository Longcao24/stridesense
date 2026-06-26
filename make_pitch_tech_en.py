#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""StrideSense Tech deck (EN, 6 pages): demo device & results → accuracy comparison → apps → hardware → next steps."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

BG=RGBColor(0x0b,0x0e,0x14); PANEL=RGBColor(0x15,0x1c,0x28); PANEL2=RGBColor(0x1b,0x23,0x30)
LINE=RGBColor(0x2a,0x34,0x45); TXT=RGBColor(0xe6,0xed,0xf3); DIM=RGBColor(0x9f,0xb0,0xc8)
ACC=RGBColor(0x4c,0xc2,0xff); GOOD=RGBColor(0x3f,0xb9,0x50); WARN=RGBColor(0xe3,0xb3,0x41)
BAD=RGBColor(0xf8,0x51,0x49); PINK=RGBColor(0xff,0x7a,0xa8); PURP=RGBColor(0xb9,0x8c,0xff)
GREY=RGBColor(0x6b,0x7a,0x90); LGN=RGBColor(0xbf,0xe0,0xc8)
FONT="PingFang SC"; NTOT=6

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]

def setfont(r,name=FONT):
    r.font.name=name; rPr=r._r.get_or_add_rPr()
    for tag in ("a:ea","a:cs"):
        el=rPr.find(qn(tag))
        if el is None: el=rPr.makeelement(qn(tag),{}); rPr.append(el)
        el.set("typeface",name)
def slide():
    s=prs.slides.add_slide(BLANK); s.background.fill.solid(); s.background.fill.fore_color.rgb=BG; return s
def rect(s,x,y,w,h,fill=None,line=None,shape=MSO_SHAPE.ROUNDED_RECTANGLE,lw=1.0):
    sp=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb=fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(lw)
    sp.shadow.inherit=False; return sp
def text(s,x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,space=3,lh=1.1):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    for i,para in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.space_after=Pt(space); p.space_before=Pt(0); p.line_spacing=lh
        for (t,sz,col,bold) in para:
            r=p.add_run(); r.text=t; r.font.size=Pt(sz); r.font.color.rgb=col; r.font.bold=bold; setfont(r)
    return tb
def pic(s,path,x,y,h=None,w=None):
    kw={}
    if h is not None: kw["height"]=Inches(h)
    if w is not None: kw["width"]=Inches(w)
    return s.shapes.add_picture(path,Inches(x),Inches(y),**kw)
def title(s,kicker,ttl,tsz=22):
    rect(s,0.55,0.42,0.10,0.55,fill=ACC,shape=MSO_SHAPE.RECTANGLE)
    text(s,0.8,0.36,12.4,0.4,[[(kicker,11.5,ACC,True)]])
    text(s,0.78,0.66,12.5,0.7,[[(ttl,tsz,TXT,True)]])
def tag(s,n): text(s,9.6,7.05,3.1,0.3,[[("StrideSense  ·  "+str(n)+"/"+str(NTOT),10.5,DIM,False)]],align=PP_ALIGN.RIGHT)
def keyline(s,txt,col=ACC,y=6.5,fs=13):
    rect(s,0.7,y,11.93,0.6,fill=PANEL2,line=col,lw=1.2)
    text(s,0.95,y,11.45,0.6,[[(txt,fs,TXT,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
def arrow(s,x,y,w,h=0.3,col=ACC):
    a=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb=col; a.line.fill.background(); a.shadow.inherit=False
def node(s,x,y,w,h,emo,label,col,es=26,ls=11):
    rect(s,x,y,w,h,fill=PANEL,line=col,lw=1.6)
    text(s,x,y+h*0.1,w,h*0.42,[[(emo,es,TXT,False)]],align=PP_ALIGN.CENTER)
    text(s,x,y+h*0.5,w,h*0.48,[[(ln,ls,TXT,True)] for ln in label.split("\n")],align=PP_ALIGN.CENTER,lh=1.04,space=0)
def note(s,txt):  # presenter speaker-notes (English bullet outline)
    s.notes_slide.notes_text_frame.text=txt

# ===================================================== 1 Cover
s=slide()
rect(s,0,5.72,13.333,0.06,fill=ACC,shape=MSO_SHAPE.RECTANGLE)
text(s,0.9,1.45,11.5,0.5,[[("Phone Demo · Specs & Accuracy Comparison",14.5,ACC,True)]])
text(s,0.85,2.05,11.7,1.2,[[("StrideSense",58,TXT,True)]])
text(s,0.9,3.4,11.8,0.6,[[("Foot-mounted 6-axis IMU · phone demo proven (feasibility)",22,DIM,False)]])
rect(s,0.9,4.14,12.0,0.54,fill=PANEL,line=WARN,lw=1.0)
text(s,1.15,4.14,11.6,0.54,[[("⚠️ Pain: L-R asymmetry · impact overload · unstable landing — invisible to the eye; training by feel = injury + plateau",12.5,TXT,True)]],anchor=MSO_ANCHOR.MIDDLE)
text(s,0.9,4.94,12.0,0.5,[[("Now: one iPhone reconstructs trajectory & gait   →   Next: commercial 6-axis IMU embedded in Fitasy shoes",15,DIM,False)]])
for i,(emo,lab) in enumerate([("📱","Live Demo"),("📊","Accuracy & Capability"),("🚀","Future Apps")]):
    x=1.0+i*4.0
    text(s,x,5.98,3.9,0.5,[[(emo+"  ",18,TXT,False),(lab,14,DIM,True)]])
text(s,0.9,6.9,11.5,0.4,[[("Xiaoyu Zhang  ·  2026",13,DIM,False)]])
tag(s,1)
note(s,"""SPEAKER NOTES · Cover  (~45s)
• Hook: "Athletes train by feel — but the foot is where the truth is, and the eye can't see it."
• Point to the ⚠ box: left-right asymmetry, impact overload, unstable landing → invisible → injury + plateau.
• What we built: today, ONE iPhone strapped to the foot already reconstructs trajectory + full gait. No app, no special hardware.
• Where it goes: a commercial 6-axis IMU pod embedded in YOUR Fitasy shoes — that is the proposal.
• Agenda: three things — live demo, accuracy & capability, future apps.
→ NEXT: "Let me show you it actually works." """)

# ===================================================== 2 Demo (video + trajectory)
s=slide(); title(s,"① Current Demo · Device & Results","Phone 6-axis IMU records real walking → reconstructs 2D / 3D trajectory",tsz=20)
pic(s,"figs/imu_sensors_en.png",2.02,1.22,w=9.3)
vy=3.78; vh=2.45; vw=vh*720/1280; ph=2.0; pw=ph*1.185; ly=3.54
tot=vw+0.5+pw+0.42+pw; vx=(13.333-tot)/2
text(s,vx-0.45,ly,vw+0.9,0.26,[[("📹 Real recording",11.5,ACC,True)]],align=PP_ALIGN.CENTER)
s.shapes.add_movie("video.mp4",Inches(vx),Inches(vy),Inches(vw),Inches(vh),poster_frame_image="figs/video_poster.jpg",mime_type="video/mp4")
arrow(s,vx+vw+0.04,vy+vh/2-0.13,0.42,0.26,col=ACC)
g2x=vx+vw+0.5
text(s,g2x,ly,pw,0.26,[[("2D trajectory · ●start→○end",11,ACC,True)]],align=PP_ALIGN.CENTER)
pic(s,"figs/fig_traj2d.png",g2x,vy,h=ph)
g3x=g2x+pw+0.42
text(s,g3x,ly,pw,0.26,[[("3D trajectory",11.5,ACC,True)]],align=PP_ALIGN.CENTER)
pic(s,"figs/fig_traj3d.png",g3x,vy,h=ph)
keyline(s,"One phone (accel + gyro) already reconstructs a full single-foot gait trajectory — an in-shoe commercial pod only makes it steadier & more complete.",col=ACC,y=6.31,fs=12)
text(s,0.7,6.95,11.9,0.3,[[("Scope: single walk ~8.6s · 6 steps · 60Hz · open path (not a loop) · step-time CV≈1.7%; shows feasibility & trajectory reconstruction, not validated vs motion-capture ground truth.",9.6,DIM,False)]])
tag(s,2)
note(s,"""SPEAKER NOTES · Demo  (~1.5 min, + live demo)
• Real data, not a mockup: an iPhone's built-in 6-axis IMU (accelerometer + gyroscope) recording a real walk.
• Left = the sensor; middle = the actual recording; right = reconstructed 2D ground path + 3D trajectory.
• Why it does not drift: every footfall is momentarily still → we zero the velocity each step (ZUPT) → no runaway drift.
• KEY LINE (say it): one phone already reconstructs a full single-foot gait trajectory; an in-shoe pod only makes it steadier & more complete.
• Be honest: one ~8.6s walk, 6 steps, 60Hz, open path — proves feasibility + reconstruction, NOT yet validated vs motion-capture.
• [LIVE DEMO — optional, ~2 min]: open the phone page, Calibrate → Walk mode → walk a few steps → show the live ground path + gait panel. Keep it short.
→ NEXT: "It works — but is it accurate enough to be useful?" """)

# ===================================================== 3 Accuracy & capability
s=slide(); title(s,"② Phone Demo vs Commercial 6-axis IMU · Accuracy & Capability","Is each feature accurate enough to serve as an absolute metric?",tsz=20)
rect(s,0.7,1.34,5.85,1.12,fill=PANEL,line=WARN,lw=1.3)
text(s,0.92,1.4,5.6,0.32,[[("📱 Phone Demo · iPhone built-in 6-axis (Bosch)",11.5,WARN,True)]])
text(s,0.92,1.76,5.6,0.62,[
  [("~60 Hz · 16-bit · fixed range · soft-strapped · weak sync",10,DIM,False)],
  [("↔ Same class of 6-axis & same algorithm — just lower rate, softer mount",10,ACC,False)],
],lh=1.16,space=2)
rect(s,6.78,1.34,5.85,1.12,fill=PANEL,line=GOOD,lw=1.3)
text(s,6.98,1.4,5.5,0.32,[[("Commercial 6-axis IMU pod (real parts)",11.5,GOOD,True)]])
pic(s,"figs/chip_imu.png",6.98,1.8,w=3.4)
text(s,10.5,1.5,2.1,1.0,[[("16-bit · ±16g",10.3,TXT,True)],[("±2000 dps",10.3,TXT,True)],[("ODR ≥1 kHz",10.3,TXT,True)],[("rigid in-shoe · hard timebase",8.8,DIM,False)]],lh=1.24,space=1)
# "enough?" summary bar: 7 absolute / 4 relative
by=2.54; bx=2.6; bw=7.8; bh=0.36; gw=bw*7/11
text(s,0.7,by,1.85,bh,[[("11 features →",10.5,TXT,True)]],align=PP_ALIGN.RIGHT,anchor=MSO_ANCHOR.MIDDLE)
rect(s,bx,by,gw,bh,fill=GOOD,line=GOOD); text(s,bx,by,gw,bh,[[("✓ 7 · absolute metrics",10.5,BG,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
rect(s,bx+gw,by,bw-gw,bh,fill=ACC,line=ACC); text(s,bx+gw,by,bw-gw,bh,[[("~ 4 · relative",10.5,BG,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
hy=3.04; rect(s,0.7,hy,11.93,0.36,fill=PANEL2,line=LINE)
for cx,cwd,t,c in [(0.9,2.5,"Feature",ACC),(3.45,2.9,"Phone demo accuracy",WARN),(6.45,2.9,"Commercial pod accuracy",GOOD),(9.62,3.0,"Capability (metric?)",ACC)]:
    text(s,cx,hy,cwd,0.36,[[(t,10.5,c,True)]],anchor=MSO_ANCHOR.MIDDLE)
rows=[("Cadence","±2–3 spm","±1–2 spm","Absolute",GOOD),
      ("Stride time","±20–40 ms","±5–15 ms","Absolute",GOOD),
      ("Gait CV","±1–1.5%","±0.5–1%","Absolute",GOOD),
      ("Ground contact GCT","±30–60 ms","±5–15 ms","Absolute",GOOD),
      ("Stance / Swing","±5–8%","±2–3%","Absolute",GOOD),
      ("Stride length","±8–15%","±2–5%","Absolute",GOOD),
      ("Gait speed","±10–20%","±3–8%","Relative",ACC),
      ("Symmetry LSI","needs both feet","±2–4% (both feet)","Absolute",GOOD),
      ("Toe clearance MTC","±3–6 cm","±1–3 cm","Relative",ACC),
      ("Foot angle","±5–10°","±1–3° (contact)","Relative",ACC),
      ("2D / 3D trajectory","high drift","higher rate + rigid","Relative",ACC)]
yy=hy+0.40; rh=0.285
for idx,(nm,d,p,cap,cc) in enumerate(rows):
    rect(s,0.7,yy,11.93,rh,fill=PANEL if idx%2==0 else PANEL2,line=None)
    text(s,0.9,yy,2.5,rh,[[(nm,10.3,TXT,True)]],anchor=MSO_ANCHOR.MIDDLE)
    text(s,3.45,yy,2.95,rh,[[(d,10,DIM,False)]],anchor=MSO_ANCHOR.MIDDLE)
    text(s,6.45,yy,2.95,rh,[[(p,10,LGN,False)]],anchor=MSO_ANCHOR.MIDDLE)
    rect(s,9.62,yy+0.03,1.62,rh-0.08,fill=cc,line=cc,lw=1.0)
    text(s,9.62,yy,1.62,rh,[[(cap,9.5,BG,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    yy+=rh
text(s,0.7,yy+0.06,11.93,0.6,[
  [("■ Absolute",9.6,GOOD,True),(" = usable directly as a metric      ",9.3,DIM,False),("■ Relative",9.6,ACC,True),(" = needs one calibration, but the change is precise.  Foot signals a watch can't measure.",9.3,DIM,False)],
  [("Source: ",9.6,WARN,True),("phone column = on-device measured, ± are estimates (single trial, uncalibrated); commercial column = datasheet / literature upper bounds, benchmarked to Vicon, Plantiga.",9.3,DIM,False)],
],lh=1.22,space=2)
tag(s,3)
note(s,"""SPEAKER NOTES · Accuracy & Capability  (~1.5 min)
• Two columns: left = the phone demo (iPhone 6-axis), right = a commercial 6-axis pod (real parts).
• Same sensor class, same algorithm — the pod just adds higher rate, rigid mount, hard timebase.
• HEADLINE (point to the bar): of 11 gait features, 7 are good enough to use as ABSOLUTE metrics; 4 are RELATIVE — one calibration, but the change is precise.
• Don't read the table line by line. Message: cadence, stride time, gait CV, ground-contact, stance/swing, stride length — all absolute. These are foot signals a watch cannot measure.
• Honesty (source row): phone ± are estimates from a single uncalibrated trial; commercial column = datasheet / literature upper bounds, benchmarked to Vicon & Plantiga.
→ NEXT: "So what do these features actually unlock?" """)

# ===================================================== 4 Future apps + detected features
s=slide(); title(s,"③ Future · What a Commercial 6-axis IMU Enables","One in-shoe pod covers both Athletic Performance + Sport Rehab")
cw=3.86; cg=0.27; cx0=(13.333-(cw*3+cg*2))/2
def appcard(x,y,emo,nm,feat,col):
    rect(s,x,y,cw,1.94,fill=PANEL,line=col,lw=1.5)
    text(s,x,y+0.12,cw,0.52,[[(emo,28,TXT,False)]],align=PP_ALIGN.CENTER)
    text(s,x,y+0.70,cw,0.36,[[(nm,11.5,col,True)]],align=PP_ALIGN.CENTER)
    rect(s,x+0.35,y+1.10,cw-0.7,0.014,fill=LINE,shape=MSO_SHAPE.RECTANGLE)
    text(s,x+0.16,y+1.16,cw-0.32,0.7,[[("Detected features",9.5,DIM,True)],[(feat,9.8,TXT,False)]],align=PP_ALIGN.CENTER,lh=1.16,space=1)
def grouplabel(y,col,txt):
    rect(s,0.6,y+0.02,0.1,0.26,fill=col,shape=MSO_SHAPE.RECTANGLE)
    text(s,0.82,y,11.8,0.3,[[(txt,12.5,col,True)]],anchor=MSO_ANCHOR.MIDDLE)
perf=[("🏃","Running form & efficiency","Cadence · GCT · flight ratio · foot angle · symmetry",WARN),
      ("🗺","Path & range of motion","2D/3D trajectory · speed · turn analysis",ACC),
      ("🏋","Training load management","Cumulative impact · steps · GCT · load asymmetry",GREY)]
rehab=[("⚖️","Return-to-sport · symmetry LSI","L-R & stride symmetry · single-leg load · stability",PINK),
       ("📉","Fatigue · landing stability","Gait CV · landing impact · cycle variability",GOOD),
       ("🦶","Trip / clearance risk alert","Toe clearance MTC · variability · foot angle",PURP)]
grouplabel(1.42,WARN,"Athletic Performance · running efficiency / training load")
for i,a in enumerate(perf): appcard(cx0+i*(cw+cg),1.74,*a)
grouplabel(3.80,PINK,"Sport Rehab · return-to-sport (ACL · ankle · overuse)")
for i,a in enumerate(rehab): appcard(cx0+i*(cw+cg),4.12,*a)
keyline(s,"Swap the report template to cover both scenarios; each app = a set of quantifiable foot features.",col=ACC,y=6.22)
tag(s,4)
note(s,"""SPEAKER NOTES · Future Apps  (~1 min)
• One pod, two scenarios — both matter to your customers.
• Athletic Performance: running form & efficiency, path / range of motion, training-load management.
• Sport Rehab — this is RETURN-TO-SPORT, not medical: symmetry / RTS readiness, fatigue & landing stability, trip / clearance risk.
• Each "app" is just a different report template over the same quantifiable foot features.
→ NEXT: "Which raises the obvious question — can this actually fit inside a shoe?" """)

# ===================================================== 5 Hardware · two in-shoe routes
s=slide(); title(s,"④ Hardware · Two Routes to In-Shoe","Complete modules are all ≥10mm — two viable in-shoe routes, no circuit design",tsz=19)
def photocard(x,y,w,photo,name,spec,col):
    h=1.98
    rect(s,x,y,w,h,fill=PANEL2,line=col,lw=1.0)
    rect(s,x+0.16,y+0.12,w-0.32,1.12,fill=RGBColor(0xee,0xf1,0xf4))
    pic(s,photo,x+w/2-0.58,y+0.18,h=1.0)
    text(s,x,y+1.28,w,0.28,[[(name,10.5,TXT,True)]],align=PP_ALIGN.CENTER)
    text(s,x,y+1.58,w,0.28,[[(spec,10,col,True)]],align=PP_ALIGN.CENTER)
# ---- Route A: complete module + custom midsole pocket ----
rect(s,0.5,1.62,6.0,4.55,fill=PANEL,line=WARN,lw=1.6)
text(s,0.72,1.74,5.6,0.32,[[("A · Fast route — off-the-shelf module + custom pocket",12.5,WARN,True)]])
text(s,0.72,2.12,5.6,0.42,[[("Complete module (BLE+battery) → custom Fitasy midsole pocket, removable",10,DIM,False)]],lh=1.12)
photocard(0.72,2.62,2.78,"figs/parts/mod_wit.jpg","WitMotion WT9011DCL","~12mm · $16–30",ACC)
photocard(3.62,2.62,2.78,"figs/parts/mod_mbient.jpg","mbientlab MMS","10mm · $130",PINK)
text(s,0.72,4.74,5.6,0.3,[[("✓ No circuit design · off-the-shelf · pilot in 1–2 wks",10.2,GOOD,True)]])
text(s,0.72,5.14,5.6,0.32,[[("✗ Thick ~10–12mm → fits a custom midsole pocket, removable",10,DIM,False)]])
text(s,0.72,5.62,5.6,0.32,[[("→ Exactly the 'buy + package' idea; custom shoe is the key",10,WARN,True)]])
# ---- Route B: tiny board + ultra-thin battery ~5mm ----
rect(s,6.83,1.62,6.0,4.55,fill=PANEL,line=GOOD,lw=1.6)
text(s,7.05,1.74,5.6,0.32,[[("B · Thin route — tiny board + ultra-thin battery ~5mm",12.5,GOOD,True)]])
text(s,7.05,2.12,5.6,0.42,[[("Tiny board (BLE+6-axis) + ultra-thin LiPo → ~5mm stack, any insole",10,DIM,False)]],lh=1.12)
photocard(7.05,2.62,2.78,"figs/parts/mod_xiao.jpg","Seeed XIAO nRF52840 Sense","21×18×3.5mm · $16",ACC)
photocard(9.95,2.62,2.78,"figs/parts/batt.png","Ultra-thin LiPo pouch","~1mm / 50mAh · $3",WARN)
text(s,7.05,4.74,5.6,0.3,[[("= ~5mm pod stack · 6-axis + BLE5.4 · built-in charging",10.2,GOOD,True)]])
text(s,7.05,5.14,5.6,0.32,[[("✓ ~5mm, thinnest & unnoticeable, fits any thin insole",10,GOOD,False)]])
text(s,7.05,5.62,5.6,0.32,[[("✗ Needs a battery connected (light assembly, no circuit design)",10,DIM,True)]])
keyline(s,"Plan: route A (fast) for the real-shoe pilot first → route B (thin) for productization; both off-the-shelf, no circuit design.",col=ACC,y=6.34,fs=11.5)
tag(s,5)
note(s,"""SPEAKER NOTES · Hardware · Two Routes  (~1.5 min)
• Honest reality first: complete off-the-shelf modules (battery + Bluetooth built in) are all ≥10mm — you can't lay one flat in a thin insole.
• Two routes, NEITHER needs us to design a circuit:
• Route A (fast): a complete module (WitMotion ~$16–30, or mbientlab) → a custom Fitasy 3D-printed midsole pocket, removable. This IS the "buy + package" idea — your custom shoe is the enabler. Pilot in 1–2 weeks.
• Route B (thin): a tiny board (Seeed XIAO, 3.5mm) + ultra-thin LiPo → a ~5mm stack that fits any insole, unnoticeable. Needs a battery connected = light assembly, still no circuit design.
• Plan: Route A for the pilot now → Route B for productization.
→ NEXT: "So here is what I'm actually proposing." """)

# ===================================================== 6 Next steps + CTA
s=slide(); title(s,"⑤ Next · From Phone Demo to In-Shoe Product","Three steps: real-shoe pilot → validate accuracy & stability → thin-pod product",tsz=19)
text(s,0.8,1.42,12.0,0.4,[[("Phone demo proves '6-axis can reconstruct gait'; commercial pod accuracy = the next target to realize in-shoe, not yet achieved.",13,DIM,False)]])
steps=[("🛒","① Off-the-shelf · pilot","Complete module (BLE+battery) in\nFitasy custom midsole pocket · removable",ACC),
       ("🎯","② Validate accuracy & stability","Confirm measurement accuracy\n& everyday reliability",GOOD),
       ("👟","③ Thin pod · product","Build a ~5mm thin pod\nseamlessly embedded in Fitasy shoes",PINK)]
sw=3.55; sgap=0.55; sx0=(13.333-(sw*3+sgap*2))/2; sy=2.2; sh=2.45
for i,(emo,ttl,desc,col) in enumerate(steps):
    x=sx0+i*(sw+sgap)
    rect(s,x,sy,sw,sh,fill=PANEL,line=col,lw=1.7)
    text(s,x,sy+0.26,sw,0.7,[[(emo,38,TXT,False)]],align=PP_ALIGN.CENTER)
    text(s,x,sy+1.2,sw,0.4,[[(ttl,12.5,col,True)]],align=PP_ALIGN.CENTER)
    text(s,x+0.15,sy+1.72,sw-0.3,0.66,[[(ln,10.5,DIM,False)] for ln in desc.split("\n")],align=PP_ALIGN.CENTER,lh=1.18,space=1)
    if i<2: arrow(s,x+sw+0.06,sy+sh/2-0.16,sgap-0.12,0.32,col=col)
cy=5.05
rect(s,0.7,cy,11.93,1.05,fill=PANEL2,line=ACC,lw=1.8)
text(s,1.0,cy+0.13,11.4,0.5,[[("🚀 Proposal: ",15,ACC,True),("start with an off-the-shelf module + custom midsole pocket; run one joint real-shoe pilot on a Fitasy model with agreed, verifiable milestones.",13,TXT,True)]])
text(s,1.0,cy+0.66,11.4,0.36,[[("Business model & pricing are in the full business deck.",11,DIM,False)]])
tag(s,6)
note(s,"""SPEAKER NOTES · Next & The Ask  (~1 min)
• Three steps: ① off-the-shelf module in a Fitasy custom midsole pocket → ② validate measurement accuracy & everyday reliability → ③ a ~5mm thin pod seamlessly embedded in your shoes.
• THE ASK (CTA box): start with an off-the-shelf module + custom pocket, and run ONE joint real-shoe pilot on a Fitasy model, with agreed, verifiable milestones. Low cost, low risk, fast.
• Business model & pricing live in the separate business deck — happy to walk through it after.
• Close: "The phone demo already proved the algorithm. The only thing left is to put it in your shoe — and that is a pilot, not a research project."
[TOTAL ≈ 7 min talk + ~2 min live demo = 9–10 min]""")

out="StrideSense_Tech_EN.pptx"; prs.save(out)
print("saved",out,"·",len(prs.slides._sldIdLst),"slides")
